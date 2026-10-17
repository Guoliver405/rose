"""
Tabelle zum Gegenlesen einer Übersetzung (17.10.2026).

    python scripts/translation-sheet.py guest en Sessions/Uebersetzung-guest-en.xlsx
    python scripts/translation-sheet.py all en Sessions/Uebersetzung-alle-en.xlsx
    python scripts/translation-sheet.py all en Sessions/Neu.xlsx --since Sessions/Uebersetzung-en-gegengelesen-2026-10-17.xlsx

Holt die Texte über `translation-export.mjs` (Sprachdateien und Erläuterungen
in `src/lib/i18n/`) und schreibt eine Excel-Datei: Schlüssel, Ort,
Erläuterung, Deutsch, Übersetzung, Korrektur, Anmerkung. Mit `all` stehen
alle Bereiche in einer Datei, je Bereich ein Reiter. Korrekturen werden von
Hand zurück in `src/lib/i18n/<sprache>/…` übertragen.

Mit `--since <gegengelesene Datei>` stehen nur Texte darin, die der Gegenleser
noch nicht gesehen hat: neue Schlüssel, geänderte deutsche Texte und
Übersetzungen, die weder seiner Vorlage noch seiner Korrektur entsprechen.
Reine Zeichensetzung (Komma vor dem letzten Glied, Punkt im Anführungszeichen)
zählt nicht als Änderung. `--full admin,…` nimmt einzelne Bereiche trotzdem
vollständig auf (die Rezeption war in der Runde vom 17.10. nicht wirklich
gelesen, die Vorschläge waren alle um Zeilen verrutscht).
"""
import re
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

LANG_NAMES = {"en": "Englisch, US", "es": "Spanisch", "fr": "Französisch"}

# Reihenfolge = Reiter-Reihenfolge bei `all`. Reitername höchstens 31 Zeichen.
AREAS = {
    "guest": ("Gästeportal", "Gästeportal"),
    "service": ("Reinigungsboard", "Reinigungsboard"),
    "admin": ("Rezeption", "Rezeption"),
    "mail": ("E-Mails", "E-Mails"),
    "print": ("Druckseiten", "Druckseiten"),
    "account": ("Konto und Abrechnung", "Konto und Abrechnung"),
    "auth": ("Anmeldung und Registrierung", "Anmeldung"),
    "help": ("Hilfe und Lotsen", "Hilfe und Lotsen"),
    "symbols": ("Zimmer-Symbole (Hilfe)", "Zimmer-Symbole"),
    "sheet": ("Check-in-Blatt und Aushang", "Check-in-Blatt"),
}

READERS = {
    "guest": "Leser ist der Gast, am Handy. Höflich und schlicht; das Haus spricht als „we“ (nie „hotel“ – auch Pensionen nutzen RoSe).",
    "service": "Leserin ist eine Reinigungskraft, am Handy zwischen zwei Zimmern. Kurz und freundlich; Begriffe wie im Gästeportal („Do not disturb“, „No cleaning today“, „front desk“). Zustände auf den Zimmer-Kacheln höchstens zwei bis drei Wörter.",
    "admin": "Leser sind Rezeption, Inhaber und Manager, am Rechner oder Tablet. Sachlich und freundlich; „property“ statt „hotel“; Begriffe wie in den Portalen („front desk“, „access card“, „check-in sheet“, „room sign“).",
    "mail": "Leser der E-Mails sind Personal und Kontoinhaber; die Statuszeile liest die Verwaltung unter einem Sende-Knopf.",
    "print": "Bedienung der Druckseiten für die Rezeption; Aufstellung und Zugangskarte werden gedruckt und ausgehändigt.",
    "account": "Leser sind Kontoinhaber (Abrechnung) bzw. Inhaber und Manager (Häuser-Seite). Steuer- und Vertragsbegriffe genau (VAT ID, reverse charge, terms and conditions). Die Rechnungstexte gehen derzeit auf Deutsch heraus.",
    "auth": "Anmeldung, Registrierung und Passwort der Verwaltung. Knapp; die Seite „Passwort zurücksetzen“ verrät bewusst nicht, ob es zu einer Adresse ein Konto gibt.",
    "help": "Leser sind Rezeption, Inhaber und Manager. Erklärend; Namen von Knöpfen und Seiten in Anführungszeichen wortgleich mit der Oberfläche. „Lotse“ heißt „guide“.",
    "symbols": "Legende der Zimmer-Kacheln in der Hilfe: kurze Namen der Zeichen und eine Erklärung je Zeichen.",
    "sheet": "Gedrucktes Blatt für den Gast und die Zugangs-E-Mail. Gibt es schon seit September viersprachig; hier zur Vollständigkeit.",
}

GENERAL = [
    "Platzhalter wie {time} oder {number} und Auszeichnungen wie <b>…</b> müssen stehen bleiben; ihre Stellung im Satz darf sich ändern.",
    "Amerikanisches Englisch (Entscheidung vom 17.10.2026).",
    "Uhrzeiten erscheinen im 24-Stunden-Format (10:00), unabhängig von der Sprache.",
    "Knopf-Beschriftungen kurz halten – manche stehen auch auf gedrucktem Papier oder auf schmalen Kacheln am Handy.",
    "Typografie: “…”, Auslassungszeichen …, Gedankenstrich — mit Leerzeichen.",
]


def export(root: Path, area: str, lang: str) -> dict:
    raw = subprocess.run(
        ["node", "--no-warnings", "--import", (root / "scripts" / "ts-resolve.mjs").as_uri(),
         str(root / "scripts" / "translation-export.mjs"), area, lang],
        check=True, capture_output=True,
    ).stdout.decode("utf-8")
    return json.loads(raw)


def norm(text) -> str:
    t = str(text or "").replace(",", "").replace(".”", "”").replace("”.", "”")
    return re.sub(r"\s+", " ", t).strip()


def reviewed(path: Path) -> dict:
    """(Reiter, Schlüssel) → (Deutsch, Vorlage, Korrektur) aus der gegengelesenen Datei."""
    from openpyxl import load_workbook
    out = {}
    for ws in load_workbook(path, read_only=True).worksheets[1:]:
        for row in list(ws.iter_rows(values_only=True))[1:]:
            if row and row[0]:
                out[(ws.title, row[0])] = (row[3], row[4], row[5])
    return out


def unseen(rows: list, sheet: str, base: dict) -> list:
    keep = []
    for r in rows:
        old = base.get((sheet, r["key"]))
        if old is None or norm(old[0]) != norm(r["de"]) or norm(r["target"]) not in {norm(old[1]), norm(old[2])}:
            keep.append(r)
    return keep


def write_texts(ws, rows: list, sprache: str, table_name: str) -> None:
    ws.append(["Schlüssel", "Ort", "Erläuterung", "Deutsch", f"Übersetzung ({sprache})", "Korrektur", "Anmerkung"])
    for r in rows:
        ws.append([r["key"], r["where"], r["note"], r["de"], r["target"], "", ""])
    for i, w in enumerate([26, 38, 48, 48, 48, 40, 30], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[0].font = Font(name="Consolas", size=9, color="666666")
        row[4].font = Font(bold=True)
        row[5].fill = PatternFill("solid", fgColor="FFF7D6")
        row[6].fill = PatternFill("solid", fgColor="FFF7D6")
    table = Table(displayName=table_name, ref=f"A1:G{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "B2"


def write_intro(ws, title: str, lines: list) -> None:
    for i, (text, bold) in enumerate([(title, True), ("", False)] + lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=14 if i == 1 else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def main() -> None:
    args = sys.argv[1:]
    since = None
    if "--since" in args:
        i = args.index("--since")
        since = Path(args[i + 1])
        del args[i:i + 2]
    full = set()
    if "--full" in args:
        i = args.index("--full")
        full = set(args[i + 1].split(","))
        del args[i:i + 2]
    area, lang, out = (args[:3] + [None, None, None])[:3]
    area = area or "guest"
    lang = lang or "en"
    root = Path(__file__).resolve().parent.parent
    out = Path(out) if out else root / "Sessions" / f"Uebersetzung-{area}-{lang}.xlsx"
    sprache = LANG_NAMES.get(lang, lang)
    areas = list(AREAS) if area == "all" else [area]

    wb = Workbook()
    intro = wb.active
    intro.title = "Anleitung"
    data = {a: export(root, a, lang) for a in areas}
    if since:
        base = reviewed(since)
        for a in [a for a in areas if a not in full]:
            data[a]["rows"] = unseen(data[a]["rows"], AREAS[a][1], base)
        areas = [a for a in areas if data[a]["rows"]]

    how = [
        ("So geht’s", True),
        ("1. Je Bereich ein Reiter. Spalte „Übersetzung“ lesen; Ort und Erläuterung sagen, wo der Text steht und was dabei zählt.", False),
        ("2. Ist etwas falsch oder holprig, die bessere Fassung in „Korrektur“ eintragen. Leer lassen heißt: passt.", False),
        ("3. Fragen oder Zweifel gehören in „Anmerkung“.", False),
        ("", False),
        ("Für alle Texte", True),
    ] + [(f"• {g}", False) for g in GENERAL]
    if area == "all":
        how += [("", False), ("Bereiche", True)]
        how += [(f"• {AREAS[a][0]} ({len(data[a]['rows'])} Texte): {READERS[a]}", False) for a in areas]
        total = sum(len(d["rows"]) for d in data.values())
        how += [("", False), (f"Zusammen {total} Texte · erzeugt aus src/lib/i18n/", False)]
        title = f"RoSe – alle Texte, Übersetzung {sprache} zum Gegenlesen"
        if since:
            title = f"RoSe – neue und geänderte Texte, Übersetzung {sprache} zum Gegenlesen"
            how += [("", False), ("Nur Texte, die seit der letzten Runde neu sind oder sich geändert haben; alles andere ist bereits gegengelesen."
                     + (f" Vollständig enthalten: {', '.join(AREAS[a][0] for a in areas if a in full)}." if full else ""), False)]
    else:
        how += [(f"• {READERS.get(area, '')}", False), ("", False),
                (f"{len(data[area]['rows'])} Texte · erzeugt aus src/lib/i18n/", False)]
        title = f"{AREAS.get(area, (area,))[0]} – Übersetzung {sprache} zum Gegenlesen"
    write_intro(intro, title, how)

    for a in areas:
        ws = wb.create_sheet(AREAS[a][1] if area == "all" else "Texte")
        write_texts(ws, data[a]["rows"], sprache, f"Texte_{a}")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{out} ({sum(len(d['rows']) for d in data.values())} Texte)")


if __name__ == "__main__":
    main()
